from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).parent


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10)
    for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11)]:
        style = doc.styles[name]
        style.font.name = 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('CareConnect | SWEN 661 | Week 6').font.size = Pt(8)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for index, title in enumerate(headers):
        hdr[index].text = title
        shade(hdr[index], '1B4F72')
        for run in hdr[index].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for c_index, text in enumerate(row):
            cells[c_index].text = text
            if r_index % 2:
                shade(cells[c_index], 'F3F7FA')
            for paragraph in cells[c_index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_border(cell)
            if widths:
                cell.width = Inches(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
    doc.add_paragraph()
    return table


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = s.add_run(subtitle)
    run.bold = True
    run.font.size = Pt(11)


criteria = [
    ('1.1.1 Non-text Content', 'Text alternatives and meaningful labels for logos, icon-only controls, status and actions.', 'Implemented; automated semantic inspection; device reader check pending.'),
    ('1.2.1–1.2.5 Time-based Media', 'No prerecorded audio/video content is presented in either app.', 'Not applicable.'),
    ('1.3.1 Info and Relationships', 'Visible form labels are programmatically exposed; roles/states identify tabs, buttons and switches.', 'Implemented; RN label/hint tests added.'),
    ('1.3.2 Meaningful Sequence', 'Logical top-to-bottom layout and native navigation order; modal content is grouped.', 'Implemented; device reader check pending.'),
    ('1.3.3 Sensory Characteristics', 'Instructions use names and text, not position/color alone.', 'Implemented; code review.'),
    ('1.3.4 Orientation', 'Apps do not block supported device orientation beyond their intended portrait presentation.', 'Review on target device before submission.'),
    ('1.3.5 Identify Input Purpose', 'Sign-in and registration inputs use recognizable labels and appropriate keyboard types.', 'Implemented; device autofill check pending.'),
    ('1.4.1 Use of Color', 'Labels, borders, icons and text accompany state/color indicators.', 'Implemented; code review.'),
    ('1.4.2 Audio Control', 'No auto-playing audio.', 'Not applicable.'),
    ('1.4.3 Contrast Minimum', 'Token palette documents AA contrast pairs for normal text and primary/destructive controls.', 'Implemented; inspect custom text on device.'),
    ('1.4.4 Resize Text', 'Settings provides a 200% in-app text-size option; native platform scaling remains enabled.', 'Implemented; automated RN scale test; device reflow check pending.'),
    ('1.4.5 Images of Text', 'No essential text is rendered as an image.', 'Implemented; code review.'),
    ('1.4.10 Reflow', 'Scrollable screen and sheet layouts avoid a required horizontal scroll path at larger text sizes.', 'Implemented; validate at 200% on devices.'),
    ('1.4.11 Non-text Contrast', 'Focus/border/error tokens provide visible boundaries for inputs and controls.', 'Implemented; code review.'),
    ('1.4.12 Text Spacing', 'Native text widgets and scrollable layouts accommodate user text settings.', 'Validate with platform accessibility settings.'),
    ('1.4.13 Content on Hover or Focus', 'No hover-only essential content; touch/keyboard controls remain operable.', 'Implemented; keyboard/device check pending.'),
    ('2.1.1 Keyboard', 'Flutter controls use FocusableActionDetector/native controls; React Native relies on platform focusable controls.', 'Validate with hardware keyboard on supported target.'),
    ('2.1.2 No Keyboard Trap', 'Native modal dismissal and navigation use explicit close/cancel actions.', 'Implemented; manual keyboard check pending.'),
    ('2.1.4 Character Key Shortcuts', 'No single-character shortcuts are implemented.', 'Not applicable.'),
    ('2.2.1 Timing Adjustable', 'Demo sign-in delay does not impose a user deadline.', 'Implemented.'),
    ('2.2.2 Pause Stop Hide', 'No moving, blinking or auto-updating essential content.', 'Not applicable.'),
    ('2.3.1 Three Flashes', 'No flashing content is used.', 'Not applicable.'),
    ('2.4.1 Bypass Blocks', 'Native tab navigation gives direct access to principal app sections.', 'Implemented.'),
    ('2.4.2 Page Titled', 'App identity and screen headings identify context.', 'Implemented; device reader check pending.'),
    ('2.4.3 Focus Order', 'Native navigation and modal controls preserve a logical order.', 'Implemented; device reader check pending.'),
    ('2.4.4 Link Purpose', 'Visible action labels describe destinations/actions.', 'Implemented; code review.'),
    ('2.4.5 Multiple Ways', 'Bottom tabs, dashboard shortcuts and list routes provide more than one path to core functions.', 'Implemented.'),
    ('2.4.6 Headings and Labels', 'Screen headings and input/action labels are descriptive.', 'Implemented; label tests added.'),
    ('2.4.7 Focus Visible', 'Flutter primary controls render a visible focus ring; native controls retain platform focus behavior.', 'Implemented; keyboard/device check pending.'),
    ('2.5.1 Pointer Gestures', 'All primary flows use single-pointer taps.', 'Implemented.'),
    ('2.5.2 Pointer Cancellation', 'Native press controls provide standard cancellation behavior.', 'Implemented.'),
    ('2.5.3 Label in Name', 'Accessible names use visible control labels.', 'Implemented; RNTL role/name tests.'),
    ('2.5.4 Motion Actuation', 'No motion-activated actions are required.', 'Not applicable.'),
    ('3.1.1 Language of Page', 'Default application language is the platform English locale.', 'Verify platform metadata before release.'),
    ('3.1.2 Language of Parts', 'No language changes are declared in the current UI.', 'Not applicable.'),
    ('3.2.1 On Focus', 'Focus does not cause unexpected navigation or submission.', 'Implemented.'),
    ('3.2.2 On Input', 'Typing does not unexpectedly change context; submissions require explicit controls.', 'Implemented.'),
    ('3.2.3 Consistent Navigation', 'The same tab navigation order and labels repeat across app screens.', 'Implemented.'),
    ('3.2.4 Consistent Identification', 'Shared actions use consistent labels and roles.', 'Implemented.'),
    ('3.3.1 Error Identification', 'Required-field errors state the missing/corrective information and are announced politely.', 'Implemented; RN error-hint tests.'),
    ('3.3.2 Labels or Instructions', 'Visible field labels, hints and example placeholders are supplied.', 'Implemented; RN label propagation tests.'),
    ('3.3.3 Error Suggestion', 'Validation text gives a correction, for example enter a dose or a six-character password.', 'Implemented.'),
    ('3.3.4 Error Prevention', 'Destructive actions and emergency dialing require confirmation.', 'Implemented; workflow/manual check pending.'),
    ('4.1.1 Parsing', 'Flutter/Dart and TypeScript are compiled through platform tooling; no web markup is delivered.', 'Not applicable to native UI; build validation required.'),
    ('4.1.2 Name Role Value', 'Interactive controls expose native roles, states, labels and values through Semantics/RN accessibility props.', 'Implemented; automated semantic/RNTL coverage.'),
    ('4.1.3 Status Messages', 'Validation and task feedback are visible and announced from live regions/semantic status.', 'Implemented; device reader check pending.'),
    ('2.5.7 Dragging Movements (WCAG 2.2 extension)', 'Dashboard reordering has Move Up and Move Down buttons; no drag-only workflow.', 'Implemented; Maestro/manual workflow check.'),
]


def create_vpat():
    doc = Document()
    style_doc(doc)
    add_title(doc, 'CareConnect Mobile Accessibility Conformance Report', 'SWEN 661 Week 6 Accessibility and UI Testing')
    doc.add_paragraph('Product: CareConnect Flutter and React Native mobile applications')
    doc.add_paragraph('Prepared for: Team 1 | Date: September 21, 2026')
    doc.add_heading('Conformance summary', level=1)
    doc.add_paragraph(
        'CareConnect implements the accessibility behaviors listed below in both mobile applications. '
        'This report distinguishes implemented and automated evidence from observations that require a real TalkBack or VoiceOver session. '
        'The final submitted copy must be updated after the device test log is completed.'
    )
    add_table(doc, ['Platform', 'Accessibility implementation', 'Evidence'], [
        ('Flutter', 'Semantics, labelled controls, 48dp primary controls, focus ring, live validation feedback, 200% text-size setting, tap-based reordering.', 'Flutter accessibility guideline test and integration workflow.'),
        ('React Native', 'Accessible roles/states, visible labels propagated to TextInput, polite validation feedback, 44dp+ controls, 200% text-size setting, tap-based reordering.', 'RNTL role/state/form tests and Maestro workflows.'),
    ], [1.1, 3.7, 2.3])
    doc.add_heading('WCAG 2.1 Level AA criteria', level=1)
    doc.add_paragraph('Status wording: “Implemented” means the behavior is present in this codebase and has code or automated-test evidence. “Pending” identifies a required final device/assistive-technology verification, not a claim that it was performed.')
    add_table(doc, ['Criterion', 'CareConnect implementation', 'Status and verification'], criteria, [1.55, 3.55, 2.0])
    doc.add_heading('Known limitations and completion actions', level=1)
    for text in [
        'TalkBack and VoiceOver behavior cannot be certified from a source-code review. Complete the four rows in SCREEN_READER_TEST_LOG.md and attach the two requested recordings.',
        'Automated coverage is complete: Flutter achieved 93.90% line coverage and React Native achieved 82.08% line coverage. Run Maestro on installed builds and save that device output with the submission evidence.',
        'Recheck text reflow, focus visibility, modal focus return, and custom contrast pairs at 200% text size on the actual Android and iOS builds.',
    ]:
        doc.add_paragraph(text, style='List Bullet')
    doc.save(OUT / 'CareConnect_Week6_VPAT.docx')


def create_submission_items():
    doc = Document()
    style_doc(doc)
    add_title(doc, 'CareConnect Week 6 Submission Items', 'SWEN 661 User Interface Implementation')
    add_table(doc, ['Team', 'Members', 'Date'], [
        ('Team 1', 'Dom Puller, Wiliss Tako, Rehman Uddin', 'September 21, 2026'),
    ], [1.0, 4.4, 1.7])
    doc.add_heading('Submission summary', level=1)
    doc.add_paragraph('This package documents the Week 6 accessibility and UI testing work for both CareConnect mobile applications. It includes implementation evidence, automated testing assets, coverage instructions, and the required device-testing record.')
    add_table(doc, ['Deliverable', 'Location or evidence', 'Final submit action'], [
        ('Accessible Flutter app', 'mobile-flutter/ with Semantics, 200% text size, guideline test and integration test.', 'Flutter tests passed; run device tests.'),
        ('Accessible React Native app', 'mobile-rn/ with labelled inputs, announced errors, 200% text size and RNTL tests.', '15 suites / 131 tests passed; run device tests.'),
        ('E2E tests', 'maestro/ contains six critical workflow definitions.', 'Execute on installed Android/iOS builds; save output.'),
        ('Coverage evidence', 'Flutter: 93.90% lines. React Native: 79.97% statements / 82.08% lines.', 'Attach the generated coverage reports.'),
        ('VPAT', 'docs/week6/CareConnect_Week6_VPAT.docx', 'Update pending device observations, then submit PDF or Word.'),
        ('Screen reader evidence', 'docs/week6/SCREEN_READER_TEST_LOG.md', 'Complete TalkBack and VoiceOver rows and attach two recordings.'),
    ], [1.3, 3.6, 2.2])
    doc.add_heading('Repository build and test commands', level=1)
    for cmd in [
        'Flutter: flutter pub get; flutter test --coverage; flutter test integration_test',
        'React Native: npm ci; npm run coverage; npm run lint',
        'Maestro: maestro test maestro/<flow>.yaml after installing a platform build',
    ]:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(cmd)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
    doc.save(OUT / 'Team1_Week6_Submission_Items.docx')


if __name__ == '__main__':
    create_vpat()
    create_submission_items()
